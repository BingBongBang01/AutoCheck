"""
검사 커맨드 카탈로그 — 점검 때 장비에 이 순서대로 입력하는 명령 목록.

화면(커맨드 카탈로그)은 메모장처럼 고친다: 한 줄에 명령 하나, 명령 다음 줄에 '!'(세션 점검이 명령 사이에 넣어 보내는 것과
같다 — api/terminal_inspection_api). 저장 파일(commands_catalog.yaml)은 항목 목록 [{id, command, enabled, check_id}] 이고,
세션 점검·무인 점검·벤더 번역(engine/vendor_commands)이 이 목록을 읽는다. 예전 화면(행마다 켜고 끄기·카테고리·설명)이
남긴 파일도 그대로 읽는다 — 꺼 둔 항목과 정책에 걸리는 항목은 메모장에 보이지 않고(점검에서 실행되지 않던 것), 다음에
저장할 때 빠진다.

check_id(기본 카탈로그 엑셀의 '체크ID' — version_info, cpu_usage …)는 '무엇을 확인하는 명령인지'다. 카탈로그 명령은
Arista EOS 로 쓰여 있고, Cisco 장비에는 같은 체크ID 의 Cisco 명령을 보낸다(engine/vendor_commands, IMPROVEMENT_PLAN
4-5). 체크ID 가 생기기 전에 저장된 카탈로그는 읽을 때 기본 카탈로그와 명령 글자가 같은 항목에 체크ID 를 잇는다.
"""
import functools
import os
import yaml
from core.atomic_io import dump_yaml_atomic
from core.command_policy import blocked_reason

DEFAULT_CATALOG_PATH = "config/commands_catalog.yaml"

CATEGORY_LABELS = {"essential": "필수", "optional": "선택사항", "custom": "커스텀"}

from core.paths import AppPaths

EXCEL_DEFAULT_PATH = str(AppPaths.bundle_root() / "config" / "commands_catalog_default.xlsx")

_TRUTHY = {"true", "1", "y", "yes", "o", "사용", "예", "checked"}


def _parse_enabled(value, default):
    """엑셀 '사용여부' 칸 -> bool. 'FALSE' 같은 문자열은 bool() 로 바꾸면 참이 되므로 직접 해석한다."""
    if isinstance(value, bool):
        return value
    if value is None or (isinstance(value, float) and value != value):  # 빈 칸(None/NaN)
        return default
    return str(value).strip().lower() in _TRUTHY


def _cell_text(value):
    return "" if value is None else str(value).strip()


@functools.lru_cache(maxsize=1)
def _default_rows():
    """기본 카탈로그 엑셀의 행 — ({열 이름: 값}, …). 빈 행도 자리를 지킨다(항목 id 의 번호가 행 번호다).

    예전엔 pandas 로 읽었다(이 한 곳에서만 쓰여 앱을 켤 때 665 ms 를 먹지 않게 늦게 불렀다). 이제 저장된 카탈로그에
    체크ID 를 이을 때도 읽으므로, 이미 쓰는 openpyxl 로 한 번만 읽어 둔다."""
    if not os.path.exists(EXCEL_DEFAULT_PATH):
        print(f"경고: 기본 커맨드 엑셀 파일이 없습니다. ({EXCEL_DEFAULT_PATH}) 빈 카탈로그로 초기화합니다.")
        return ()
    try:
        import openpyxl
        workbook = openpyxl.load_workbook(EXCEL_DEFAULT_PATH, read_only=True, data_only=True)
        try:
            rows = workbook.active.iter_rows(values_only=True)
            headers = [_cell_text(cell) for cell in next(rows, ())]
            return tuple(dict(zip(headers, row)) for row in rows)
        finally:
            workbook.close()
    except Exception as e:
        print(f"엑셀 파일 읽기 실패: {e}")
        return ()


def _make_default_catalog():
    catalog = []
    for idx, row in enumerate(_default_rows()):
        command_text = _cell_text(row.get("커맨드"))
        if not command_text:
            continue
        raw_category = _cell_text(row.get("카테고리"))
        category = "essential" if raw_category == "필수" else ("optional" if raw_category == "선택사항" else "custom")
        item = {
            "id": f"{category}_{idx}",
            "category": category,
            "command": command_text,
            "description": _cell_text(row.get("설명")),
            "enabled": _parse_enabled(row.get("사용여부"), default=True),
        }
        if _cell_text(row.get("체크ID")):
            item["check_id"] = _cell_text(row.get("체크ID"))
        catalog.append(item)
    return catalog


def _command_key(command):
    return " ".join(str(command or "").split()).lower()


@functools.lru_cache(maxsize=1)
def _default_check_ids():
    return {_command_key(item["command"]): item["check_id"] for item in _make_default_catalog() if item.get("check_id")}


def with_check_ids(catalog):
    """체크ID 가 없는 항목에 기본 카탈로그에서 명령 글자가 같은 항목의 체크ID 를 잇는다(제자리, catalog 를 돌려준다).
    체크ID 가 생기기 전에 저장된 카탈로그용 — 다음에 저장할 때 파일에도 남는다. 기본에 없는 명령(사용자가 더했거나
    글자를 고친 것)은 체크ID 없이 둔다: 모든 장비에 그 글자 그대로 보낸다."""
    known = _default_check_ids()
    for item in catalog:
        if not item.get("check_id") and _command_key(item.get("command")) in known:
            item["check_id"] = known[_command_key(item.get("command"))]
    return catalog


def load_catalog(path=DEFAULT_CATALOG_PATH):
    if not os.path.exists(path):
        catalog = _make_default_catalog()
        save_catalog(catalog, path)
        return catalog
    with open(path, encoding="utf-8") as f:
        return with_check_ids(yaml.safe_load(f) or [])


def save_catalog(catalog, path=DEFAULT_CATALOG_PATH):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    # "blocked" 는 화면 표시용으로 붙이는 계산값이다(api/catalog_api.get_catalog) — 저장하지 않는다.
    dump_yaml_atomic([{k: v for k, v in item.items() if k != "blocked"} for item in catalog], path)


def split_enabled_commands(catalog):
    """활성 항목을 (실행할 명령 목록, [(실행하지 않을 명령, 이유)]) 로 나눈다.

    읽기 전용 정책(core/command_policy.py)에 걸리는 명령은 켜져 있어도 실행하지 않는다 —
    v0.5.9 이전 기본 카탈로그를 받은 사용자 파일에는 `conf` 가 켜진 채 남아 있다."""
    allowed, skipped = [], []
    for item in catalog:
        if not item.get("enabled"):
            continue
        command = item["command"]
        reason = blocked_reason(command)
        if reason:
            skipped.append((command, reason))
        else:
            allowed.append(command)
    return allowed, skipped


def get_enabled_commands(catalog):
    """실행할 명령 — 활성 항목 중 읽기 전용 정책을 통과한 literal command 그대로(벤더 무관)."""
    return split_enabled_commands(catalog)[0]


def import_from_excel(path):
    """예전 화면이 내보낸 엑셀(열: 순서, 카테고리, 사용여부, 커맨드, 설명[, 체크ID])을 읽어 catalog 리스트로 반환 —
    불러오기가 .txt 와 함께 받는다(read_catalog_file). 체크ID 열이 없는 예전 엑셀은 기본 카탈로그와 명령 글자가 같은 항목에
    체크ID 를 잇는다(with_check_ids)."""
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("openpyxl 미설치 — pip install openpyxl 후 다시 시도하세요.")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    label_to_category = {v: k for k, v in CATEGORY_LABELS.items()}

    parsed = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[3]:
            continue
        order, category_label, enabled_raw, command, description, check_id = (list(row) + [None] * 6)[:6]
        category = label_to_category.get(str(category_label).strip(), "custom") if category_label else "custom"
        parsed.append({
            "order": order if isinstance(order, (int, float)) else len(parsed) + 1,
            "category": category,
            "enabled": _parse_enabled(enabled_raw, default=False),
            "command": str(command).strip(),
            "description": str(description).strip() if description else "",
            "check_id": _cell_text(check_id),
        })
    parsed.sort(key=lambda x: x["order"])

    counters = {"essential": 0, "optional": 0, "custom": 0}
    catalog = []
    for item in parsed:
        cat = item["category"]
        idx = counters[cat]
        counters[cat] += 1
        entry = {
            "id": f"{cat}_{idx}",
            "category": cat,
            "command": item["command"],
            "description": item["description"] or "(설명 없음)",
            "enabled": item["enabled"],
        }
        if item["check_id"]:
            entry["check_id"] = item["check_id"]
        catalog.append(entry)
    return with_check_ids(catalog)


# ---------- 메모장 편집(커맨드 카탈로그 화면) ----------
SEPARATOR = "!"
# 불러오기가 받는 파일 크기 — 명령 수십 개는 몇 KB 다. 로그 파일을 잘못 고르면 화면이 통째로 멈추지 않게 막는다.
MAX_CATALOG_FILE_BYTES = 512 * 1024


def parse_catalog_text(text):
    """메모장 글 -> [(줄 번호, 명령)]. 빈 줄과 '!' 로 시작하는 줄(명령 사이 구분선·CLI 주석)은 명령이 아니다. 앞뒤 공백은 뗀다."""
    commands = []
    for number, line in enumerate(str(text or "").splitlines(), start=1):
        command = line.strip()
        if command and not command.startswith(SEPARATOR):
            commands.append((number, command))
    return commands


def commands_to_text(commands, newline="\n"):
    """명령 목록 -> 메모장 글: 명령마다 그다음 줄에 '!'."""
    return "".join(f"{command}{newline}{SEPARATOR}{newline}" for command in commands)


def shown_commands(catalog):
    """메모장에 보일 명령 — 켜져 있고 읽기 전용 정책을 통과한 것(점검에서 실제로 실행되는 명령)."""
    return [item["command"] for item in catalog if item.get("enabled") and not blocked_reason(item["command"])]


def hidden_items(catalog):
    """메모장에 보이지 않는 예전 항목 [(명령, 이유)] — 꺼 둔 것, 정책에 걸리는 것. 다음에 저장하면 카탈로그에서 빠진다."""
    hidden = []
    for item in catalog:
        if not item.get("enabled"):
            hidden.append((item["command"], "꺼 둔 명령"))
        elif blocked_reason(item["command"]):
            hidden.append((item["command"], blocked_reason(item["command"])))
    return hidden


def check_catalog_text(text):
    """메모장 글을 저장해도 되는지 — (명령 목록, 막힌 줄, 확인할 줄). 줄은 {"line", "command", "reason"} — 확인할 줄에는
    명령 목록의 순번(index)도 있다(저장하며 줄을 정리하면 그 명령은 2*index+1 번째 줄이 된다).

    막힌 줄: 장비 설정을 바꿀 수 있는 명령(읽기 전용 정책 core/command_policy) — 하나라도 있으면 저장하지 않는다.
    확인할 줄(저장은 한다): 한글처럼 ASCII 가 아닌 글자가 든 명령 — 한/영 전환 실수일 수 있다(예전 기본 카탈로그를 받은
    고객사 파일에 'show spanning-tree 피무'(vlan)가 남아 있다), 앞에 나온 명령과 같은 명령."""
    commands, blocked, warnings, seen = [], [], [], set()
    for number, command in parse_catalog_text(text):
        reason = blocked_reason(command)
        if reason:
            blocked.append({"line": number, "command": command, "reason": reason})
            continue
        if not command.isascii():
            warnings.append({"line": number, "index": len(commands), "command": command,
                             "reason": "한글 등 영문이 아닌 글자가 있습니다 — 한/영 전환 실수가 아닌지 확인하세요(장비가 받지 않습니다)"})
        key = _command_key(command)
        if key in seen:
            warnings.append({"line": number, "index": len(commands), "command": command,
                             "reason": "앞에 같은 명령이 있습니다 — 장비에 두 번 보냅니다"})
        seen.add(key)
        commands.append(command)
    return commands, blocked, warnings


def catalog_from_commands(commands, previous=()):
    """명령 목록 -> 카탈로그 항목(모두 켬, 적은 순서대로). 체크ID 는 글자가 같은 이전 항목의 것을 잇고, 없으면 기본
    카탈로그에서 찾는다(with_check_ids) — 이전 화면에서 명령 글자를 고쳐도 지켜 온 체크ID 를 잃지 않게."""
    known = {_command_key(item.get("command")): item["check_id"] for item in previous if item.get("check_id")}
    catalog = []
    for index, command in enumerate(commands):
        item = {"id": f"cmd_{index}", "command": command, "enabled": True}
        if _command_key(command) in known:
            item["check_id"] = known[_command_key(command)]
        catalog.append(item)
    return with_check_ids(catalog)


def default_commands():
    """기본 카탈로그(엑셀)의 명령 [{command, description}] — 화면이 '지금 글에 없는 기본 명령'을 눌러 더하게 보여 준다."""
    return [{"command": item["command"], "description": item.get("description", "")}
            for item in _make_default_catalog() if not blocked_reason(item["command"])]


def default_catalog():
    """기본 카탈로그(엑셀)를 저장 형식의 항목으로 — '기본값으로 초기화'. 정책에 걸리는 명령은 넣지 않는다."""
    defaults = _make_default_catalog()
    return catalog_from_commands(shown_commands(defaults), defaults)


def read_catalog_file(path):
    """불러오기 — .txt(메모장 글, 내보내기 형식) 또는 예전 화면이 내보낸 .xlsx 의 명령 목록(켜진 것만, 적힌 순서).
    .txt 는 UTF-8(BOM 포함) 이 아니면 한글 Windows 메모장의 ANSI(cp949)로 읽는다."""
    if os.path.getsize(path) > MAX_CATALOG_FILE_BYTES:
        raise ValueError(f"카탈로그 파일이 아닌 것 같습니다 — {MAX_CATALOG_FILE_BYTES // 1024}KB 보다 큽니다")
    if str(path).lower().endswith(".xlsx"):
        return [item["command"] for item in import_from_excel(path) if item.get("enabled")]
    with open(path, "rb") as f:
        raw = f.read()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp949", errors="replace")
    return [command for _number, command in parse_catalog_text(text)]
