"""
범용 정기점검 보고서 엑셀 렌더러 (openpyxl).

첨부된 3개 레거시 보고서(대한전선 Arista / LGES 오창 / NC US1 특별점검)의 공통 구조를
"고객사·벤더에 종속되지 않는" 하나의 워크북 레이아웃으로 통합한 것이다:

    표지        고객사 + 프로파일(회차) 이름으로 제목을 조립 + 담당자/점검자/점검일자 서명란
    장비현황     Hostname / 모델 / IP / S/N / OS / 용도 / 비고 (장비 수만큼 동적으로 팽창)
    점검요약     장비별 특이사항 + 조치 및 점검 소견 (NC 보고서의 '0-1. 요약' 시트 대응)
    <Hostname>  장비 1대당 1시트 — 기본정보 헤더 + 점검항목 표(기준값/전월/당월) + 특이사항

그리기만 한다 — 무엇을 그릴지(장비별 항목 판정)는 engine/inspection_evaluator.py, 데이터 조립·저장은
engine/inspection_report_builder.py 가 맡는다(판정 코드가 예전엔 이 파일에 있었다 — IMPROVEMENT_PLAN 3-6).
"""
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# 판정 상태 문자열은 openpyxl 과 무관하므로 report/inspection_status.py 로 내렸다.
# 여기서 재노출해 기존 import 경로(`from report.inspection_excel import STATUS_OK`)를 지킨다.
from report.inspection_status import (
    STATUS_NA, STATUS_OK, STATUS_SKIP, STATUS_UNREACHABLE, STATUS_WARN,
)

__all__ = [
    "build_workbook",
    "STATUS_OK", "STATUS_WARN", "STATUS_NA", "STATUS_SKIP", "STATUS_UNREACHABLE",
]

_EXCEL_FORBIDDEN_SHEET_CHARS = re.compile(r"[:\\/?*\[\]]")


# --------------------------------------------------------------------------- 서식

_THIN = Side(style="thin", color="9AA0A6")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
FILL_HEAD = PatternFill("solid", start_color="D9E1F2")
FILL_LABEL = PatternFill("solid", start_color="F2F2F2")
FILL_WARN = PatternFill("solid", start_color="FFF2CC")
FILL_BAD = PatternFill("solid", start_color="FCE4E4")
FONT_BAD = Font(bold=True, color="C00000", size=10)
FONT_HEAD = Font(bold=True, size=10)
FONT_BODY = Font(size=10)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
TOP_LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)


def _put(ws, coord, value, *, font=FONT_BODY, fill=None, align=CENTER, border=BORDER,
          number_format=None):
    """셀 하나에 값 + 서식을 넣는다. 서식을 셀별로 새로 만들지 않고 모듈 상수를 공유하므로
    수백 셀을 채워도 스타일 객체가 늘어나지 않는다."""
    cell = ws[coord]
    cell.value = value
    cell.font = font
    cell.alignment = align
    if fill is not None:
        cell.fill = fill
    if border is not None:
        cell.border = border
    if number_format:
        cell.number_format = number_format
    return cell


def _merge(ws, ref, value, **kwargs):
    """병합 셀에 값을 넣는다 — 병합 범위의 왼쪽 위 셀에만 값이 들어가지만, 테두리는
    범위 전체에 그려야 표가 끊겨 보이지 않으므로 모든 셀에 border를 적용한다."""
    ws.merge_cells(ref)
    first = ref.split(":")[0]
    _put(ws, first, value, **kwargs)
    for row in ws[ref]:
        for cell in row:
            cell.border = kwargs.get("border", BORDER)
    return ws[first]


def _widths(ws, mapping):
    for column, width in mapping.items():
        ws.column_dimensions[column].width = width


def _status_style(status):
    """점검 결과 상태에 맞는 (폰트, 배경) — 접속 불가/확인필요는 눈에 바로 띄게."""
    if status == STATUS_UNREACHABLE:
        return FONT_BAD, FILL_BAD
    if status == STATUS_WARN:
        return Font(bold=True, size=10, color="7F6000"), FILL_WARN
    if status in (STATUS_NA, STATUS_SKIP):
        # 미수집/해당없음은 '판정하지 않음'이라 회색으로 죽여 둔다 — 정상(검정)과 같은 무게로
        # 보이면 점검했다는 착각을 준다.
        return Font(size=10, color="808080"), None
    return FONT_BODY, None


def _safe_sheet_title(name, used):
    """엑셀 시트명 제약(31자, : \\ / ? * [ ] 금지)에 맞추고 중복이면 _2, _3을 붙인다."""
    title = _EXCEL_FORBIDDEN_SHEET_CHARS.sub("_", str(name or "장비")).strip() or "장비"
    title = title[:31]
    if title.lower() not in used:
        used.add(title.lower())
        return title
    for suffix in range(2, 100):
        candidate = f"{title[:31 - len(str(suffix)) - 1]}_{suffix}"
        if candidate.lower() not in used:
            used.add(candidate.lower())
            return candidate
    used.add(title.lower())
    return title


# --------------------------------------------------------------------------- 시트 빌더

def _build_cover(ws, ctx):
    """표지 — 제목은 요구사항대로 고객사명 + 프로파일(회차)명으로 조립한다."""
    meta = ctx["template"]["meta"]
    ws.sheet_view.showGridLines = False
    _widths(ws, {"A": 3, "B": 14, "C": 16, "D": 22, "E": 22, "F": 18, "G": 3})
    for row in range(4, 12):
        ws.row_dimensions[row].height = 34

    title = f"{ctx['customer']}\n{ctx['profile']}\n({meta.get('title_suffix', '정기점검 보고서')})"
    _merge(ws, "B4:F11", title, font=Font(bold=True, size=24), align=CENTER, border=None)

    _merge(ws, "B14:F14", f"점검 대상 {len(ctx['devices'])}대  |  생성 시각 {ctx['generated_at']}",
           font=Font(size=11, color="595959"), align=CENTER, border=None)

    rows = [
        ("담당자", ctx["manager"]),
        ("점검자", ctx["inspector"]),
    ]
    row = 20
    for label, person in rows:
        _merge(ws, f"B{row}:B{row + 1}", label, font=FONT_HEAD, fill=FILL_LABEL)
        for column, header in zip("CDE", ("회사명", "담당자명", "연락처")):
            _put(ws, f"{column}{row}", header, font=FONT_HEAD, fill=FILL_HEAD)
        _put(ws, f"F{row}", "확인", font=FONT_HEAD, fill=FILL_HEAD)
        _put(ws, f"C{row + 1}", person.get("company", ""))
        _put(ws, f"D{row + 1}", person.get("name", ""))
        _put(ws, f"E{row + 1}", person.get("contact", ""))
        _put(ws, f"F{row + 1}", "")
        ws.row_dimensions[row + 1].height = 28
        row += 3

    _put(ws, f"B{row}", "점검일자", font=FONT_HEAD, fill=FILL_LABEL)
    _merge(ws, f"C{row}:F{row}", ctx["inspection_date"], font=Font(bold=True, size=11))
    ws.row_dimensions[row].height = 26


def _build_inventory_sheet(ws, ctx):
    """장비현황 — 장비 수만큼 아래로 동적으로 팽창하는 2차원 매트릭스."""
    headers = ["Hostname", "모델명", "IP", "S/N", "OS Version", "용도/역할", "위치",
               "Warranty", "비고"]
    _widths(ws, {"A": 26, "B": 22, "C": 16, "D": 18, "E": 14, "F": 16, "G": 16, "H": 14,
                  "I": 30})
    _merge(ws, f"A1:{get_column_letter(len(headers))}1", "장비현황",
           font=Font(bold=True, size=13), fill=FILL_HEAD)
    for index, header in enumerate(headers, start=1):
        _put(ws, f"{get_column_letter(index)}2", header, font=FONT_HEAD, fill=FILL_HEAD)

    for offset, device in enumerate(ctx["devices"]):
        row = 3 + offset
        values = [device["name"], device.get("model", ""), device.get("ip", ""),
                  device.get("serial", ""), device.get("os_version", ""),
                  device.get("role", ""), device.get("location", ""),
                  device.get("warranty", ""), device.get("memo", "")]
        unreachable = device.get("unreachable")
        for index, value in enumerate(values, start=1):
            _put(ws, f"{get_column_letter(index)}{row}", value,
                 font=FONT_BAD if unreachable else FONT_BODY,
                 fill=FILL_BAD if unreachable else None,
                 align=LEFT if index in (1, len(values)) else CENTER)
        if unreachable:
            ws[f"{get_column_letter(len(values))}{row}"].value = STATUS_UNREACHABLE
    ws.freeze_panes = "A3"


def _build_summary_sheet(ws, ctx):
    """점검요약 — 장비별 특이사항과 소견. NC 보고서의 '요약' 시트가 이 역할을 한다."""
    # 항목 수를 정상/비정상/미수집 세 열로 나눠 적는다 — 합계 하나만 있으면 '판정하지 못한
    # 항목'이 정상 쪽에 묻혀 점검 커버리지가 실제보다 좋아 보인다.
    headers = ["Hostname", "용도/역할", "IP", "점검결과", "점검항목", "정상", "비정상",
               "미수집/해당없음", "특이사항", "조치 및 점검 소견"]
    _widths(ws, {"A": 26, "B": 16, "C": 16, "D": 12, "E": 10, "F": 8, "G": 8, "H": 15,
                  "I": 52, "J": 30})
    _merge(ws, f"A1:{get_column_letter(len(headers))}1",
           f"{ctx['customer']} {ctx['profile']} 점검 요약",
           font=Font(bold=True, size=13), fill=FILL_HEAD)
    for index, header in enumerate(headers, start=1):
        _put(ws, f"{get_column_letter(index)}2", header, font=FONT_HEAD, fill=FILL_HEAD)

    for offset, device in enumerate(ctx["devices"]):
        row = 3 + offset
        ws.row_dimensions[row].height = 30
        overall = device.get("overall_status", STATUS_OK)
        font, fill = _status_style(overall)
        counts = device.get("status_counts") or {}
        _put(ws, f"A{row}", device["name"], align=LEFT)
        _put(ws, f"B{row}", device.get("role", ""))
        _put(ws, f"C{row}", device.get("ip", ""))
        _put(ws, f"D{row}", overall, font=font, fill=fill)
        _put(ws, f"E{row}", counts.get("total", len(device.get("items", []))))
        _put(ws, f"F{row}", counts.get(STATUS_OK, 0))
        _put(ws, f"G{row}", counts.get(STATUS_WARN, 0),
             font=Font(bold=True, size=10, color="7F6000") if counts.get(STATUS_WARN) else FONT_BODY)
        _put(ws, f"H{row}", counts.get("not_judged", 0), font=Font(size=10, color="808080"))
        _put(ws, f"I{row}", device.get("remarks") or "특이사항 없음", align=TOP_LEFT)
        _put(ws, f"J{row}", device.get("opinion", ""), align=TOP_LEFT)
    ws.freeze_panes = "A3"


def _build_device_sheet(ws, ctx, device):
    """장비 1대 상세 시트 — 기본정보(2행) + 점검항목 표 + 특이사항 블록.
    대한전선 보고서의 장비 시트 레이아웃을 그대로 따르되, 항목 수가 템플릿에 따라
    달라지므로 표의 높이는 렌더링 시점에 계산한다."""
    _widths(ws, {"A": 6, "B": 12, "C": 20, "D": 12, "E": 12, "F": 14, "G": 14,
                  "H": 26, "I": 24, "J": 24})

    # ---- 기본정보 (1~2행)
    _merge(ws, "A1:B2", "기본\n정보", font=FONT_HEAD, fill=FILL_HEAD)
    basics = [
        [("고객명", ctx["customer"]), ("점검일시", ctx["inspection_date"]),
         ("장비명/Hostname", device["name"]), ("IP", device.get("ip", ""))],
        [("용도/역할", device.get("role", "")), ("모델명", device.get("model", "")),
         ("S/N", device.get("serial", "")), ("OS Version", device.get("os_version", ""))],
    ]
    for row_index, pairs in enumerate(basics, start=1):
        for column_index, (label, value) in enumerate(pairs):
            label_col = get_column_letter(3 + column_index * 2)
            value_col = get_column_letter(4 + column_index * 2)
            _put(ws, f"{label_col}{row_index}", label, font=FONT_HEAD, fill=FILL_LABEL)
            _put(ws, f"{value_col}{row_index}", value)

    # ---- 점검항목 표 헤더 (3행)
    _put(ws, "A3", "No", font=FONT_HEAD, fill=FILL_HEAD)
    _put(ws, "B3", "구분", font=FONT_HEAD, fill=FILL_HEAD)
    _merge(ws, "C3:D3", "점검항목", font=FONT_HEAD, fill=FILL_HEAD)
    _merge(ws, "E3:G3", "점검 방법", font=FONT_HEAD, fill=FILL_HEAD)
    _put(ws, "H3", "기준값", font=FONT_HEAD, fill=FILL_HEAD)
    _put(ws, "I3", "전월 점검값", font=FONT_HEAD, fill=FILL_HEAD)
    _put(ws, "J3", "당월 점검값", font=FONT_HEAD, fill=FILL_HEAD)

    # ---- 점검항목 행
    first_row = 4
    rows = device.get("items", [])
    for offset, item in enumerate(rows):
        row = first_row + offset
        ws.row_dimensions[row].height = 26
        _put(ws, f"A{row}", item.get("no"))
        _put(ws, f"B{row}", item.get("group", ""))
        _merge(ws, f"C{row}:D{row}", item.get("name", ""), align=LEFT)
        _merge(ws, f"E{row}:G{row}", item.get("method", ""), align=LEFT,
               font=Font(size=9, color="404040"))
        _put(ws, f"H{row}", item.get("criteria", ""), align=LEFT, font=Font(size=9))
        _put(ws, f"I{row}", item.get("previous", ""), align=LEFT)
        font, fill = _status_style(item.get("status"))
        value = item.get("number") if item.get("number") is not None else item.get("value")
        _put(ws, f"J{row}", value, align=LEFT, font=font, fill=fill,
             number_format=item.get("number_format"))

    # ---- '구분' 열의 연속 동일 값 병합 (레거시 보고서의 세로 병합 표현)
    start = first_row
    for offset in range(1, len(rows) + 1):
        row = first_row + offset
        current = rows[offset]["group"] if offset < len(rows) else None
        if current != rows[offset - 1]["group"]:
            if row - 1 > start:
                _merge(ws, f"B{start}:B{row - 1}", rows[offset - 1]["group"],
                       font=FONT_HEAD, fill=FILL_LABEL)
            else:
                _put(ws, f"B{start}", rows[offset - 1]["group"], font=FONT_HEAD, fill=FILL_LABEL)
            start = row

    # ---- 특이사항 블록
    remark_row = first_row + len(rows)
    remark_end = remark_row + 3
    _merge(ws, f"A{remark_row}:B{remark_end}", "특이\n사항", font=FONT_HEAD, fill=FILL_HEAD)
    _merge(ws, f"C{remark_row}:J{remark_end}", device.get("remarks_detail") or "특이사항 없음",
           align=TOP_LEFT, font=Font(size=9))
    for row in range(remark_row, remark_end + 1):
        ws.row_dimensions[row].height = 30

    if device.get("unreachable"):
        _merge(ws, f"C{remark_row}:J{remark_end}",
               f"{STATUS_UNREACHABLE} — 원본로그가 없어 점검값을 수집하지 못했습니다.",
               align=TOP_LEFT, font=FONT_BAD)

    ws.freeze_panes = "A4"
    # 인쇄(=보고서 '페이지') 설정: 장비 1대 = 가로 1페이지. fitToPage는 sheet_properties의
    # pageSetUpPr에 있고 이 객체가 기본 None이라 새로 만들어 넣어야 한다.
    from openpyxl.worksheet.properties import PageSetupProperties
    ws.print_area = f"A1:J{remark_end}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "1:3"


# --------------------------------------------------------------------------- 엔트리포인트

def _build_audit_sheet(ws, ctx, audit):
    """실시간감시이력 — 이 회차에 실시간 감시가 남긴 경고와 처리(IMPROVEMENT_PLAN 4-8, engine/realtime_audit)."""
    headers = ["발생", "장비", "등급", "내용", "근거 줄", "처리", "처리 시각", "지속(초)", "메모"]
    _widths(ws, {"A": 16, "B": 18, "C": 8, "D": 50, "E": 50, "F": 12, "G": 16, "H": 10, "I": 40})
    _merge(ws, f"A1:{get_column_letter(len(headers))}1",
           f"{ctx['customer']} {ctx['profile']} 실시간 감시 이력", font=Font(bold=True, size=13), fill=FILL_HEAD)
    counts = audit["counts"]
    summary = (f"경고 {counts['total']}건(심각 {counts['CRITICAL']} · 주요 {counts['MAJOR']} · 경고 {counts['WARNING']})"
               f" / 자동 복구 {counts['auto']} · 수동 해결 {counts['manual']} · 무시 {counts['ignored']} · "
               f"미해결 {counts['open']} / 감시 장비 {audit['watched']}대(syslog 를 받은 장비 {audit['syslog_devices']}대)")
    if audit.get("excluded"):
        summary += f" / 보고서에 넣지 않은 장비의 경고 {audit['excluded']}건 제외"
    if audit.get("open_faults"):
        summary += " / 작성 시점에 열려 있던 장애: " + ", ".join(
            f"{f['device']} {f['component_id']} {f['label']}({f['since']}부터)" for f in audit["open_faults"])
    _merge(ws, f"A2:{get_column_letter(len(headers))}2", summary, align=LEFT)
    ws.row_dimensions[2].height = 30
    for index, header in enumerate(headers, start=1):
        _put(ws, f"{get_column_letter(index)}3", header, font=FONT_HEAD, fill=FILL_HEAD)
    for offset, row in enumerate(audit["rows"]):
        line = 4 + offset
        severity_font = FONT_BAD if row["severity"] == "CRITICAL" else FONT_BODY
        outcome_font = FONT_BAD if row["outcome"] == "open" else FONT_BODY
        values = [row["at"], row["device"], row["severity_label"],
                  row["message"] + (" (감시 시작 전 기록)" if row["earlier"] else ""), row["raw"],
                  row["outcome_label"], row["closed_at"], row["duration_sec"], row["note"]]
        for index, value in enumerate(values, start=1):
            _put(ws, f"{get_column_letter(index)}{line}", value,
                 font=severity_font if index == 3 else outcome_font if index == 6 else FONT_BODY,
                 fill=FILL_BAD if index == 6 and row["outcome"] == "open" else None,
                 align=LEFT if index in (4, 5, 9) else CENTER)
    if not audit["rows"]:
        _merge(ws, f"A4:{get_column_letter(len(headers))}4", "감시하는 동안 경고가 없었습니다.", align=LEFT)
    ws.freeze_panes = "A4"


def build_workbook(ctx: dict, audit: dict = None) -> Workbook:
    """보고서 컨텍스트(engine.inspection_report_builder.build_context()의 반환값)로
    완성된 Workbook 객체를 만든다. 저장은 호출부 책임(StorageService/파일 경로 정책 분리).
    audit(engine/realtime_audit)가 있으면 마지막에 '실시간감시이력' 시트를 붙인다."""
    wb = Workbook()
    _build_cover(wb.active, ctx)
    wb.active.title = "표지"
    _build_inventory_sheet(wb.create_sheet("장비현황"), ctx)
    _build_summary_sheet(wb.create_sheet("점검요약"), ctx)

    used_titles = {"표지", "장비현황", "점검요약", "실시간감시이력"}
    for device in ctx["devices"]:
        title = _safe_sheet_title(device["name"], used_titles)
        _build_device_sheet(wb.create_sheet(title), ctx, device)
    if audit:
        _build_audit_sheet(wb.create_sheet("실시간감시이력"), ctx, audit)
    return wb
