"""실시간 감시 이력(감사 기록) 보고서 부록 — IMPROVEMENT_PLAN 4-8d.

  1. 회차(프로파일)의 실시간 감시 저장본에서 경고마다 언제 났고 어떻게 닫혔는지(자동 복구·수동 해결·무시·미해결)를
     오래된 순으로 만든다. 장비에 남은 과거 기록(show logging)과 보고서에 넣지 않은 장비의 것은 뺀다.
  2. 비밀값(비밀번호 해시 등)은 가린다. 저장본은 읽기만 한다 — 깨진 파일을 옮기지 않고 직전 저장본(.bak)을 읽는다.
  3. 정기점검 Excel 에 '실시간감시이력' 시트, PDF 에 '부록. 실시간 감시 이력'. 부록을 못 만들어도 보고서는 만든다.
  4. 보고서를 만들기 전에 감시가 묶어 둔(10초) 변경을 저장한다 — 감시를 쓰지 않았으면 서비스를 만들지 않는다.
"""
import json
import time
from pathlib import Path

import pytest

from test_round_compare import eos_log, root, write_log  # noqa: F401 — root 는 fixture

SYSLOG_DOWN = "%LINEPROTO-5-UPDOWN: Line protocol on Interface Ethernet1, changed state to down"
SECRET_HASH = "$1$abcd$Xyz123456"
ALERTS = [
    {"alert_id": "a5", "device": "Core1", "severity": "CRITICAL", "type": "LINK_DOWN", "message": "과거 기록",
     "raw_line": SYSLOG_DOWN, "at": "2026-11-15T09:00:00", "past_record": True, "history": True},
    {"alert_id": "a1", "device": "Core1", "severity": "CRITICAL", "type": "LINK_DOWN",
     "message": "인터페이스 Ethernet1 Line protocol DOWN 감지!", "raw_line": SYSLOG_DOWN, "at": "2026-11-15T10:00:00",
     "resolved": True, "resolved_by": "Core1(config-if-Et1)#no shutdown", "resolved_at": "2026-11-15T10:02:00",
     "duration_sec": 120.0},
    {"alert_id": "a3", "device": "Core1", "severity": "WARNING", "type": "CONFIG_CHANGED", "message": "계정 변경",
     "raw_line": f"Core1(config)#username admin secret 5 {SECRET_HASH}", "at": "2026-11-15T10:10:00",
     "ignored": True, "ignored_ts": "2026-11-15 10:11:00", "ignored_note": "계획된 작업"},
    {"alert_id": "a2", "device": "Agg1", "severity": "MAJOR", "type": "keyword_inactive",
     "message": "키워드 'INACTIVE' 검출", "raw_line": "state : Inactive", "at": "2026-11-15T10:05:00",
     "resolved": True, "resolved_by": "현장 확인 후 정상", "resolved_at": "2026-11-15T10:20:00", "duration_sec": None},
    {"alert_id": "a4", "device": "Core1", "severity": "CRITICAL", "type": "MLAG_PEER_DOWN",
     "message": "MLAG peer down", "raw_line": "%MLAG-3-PEER_DOWN: peer down", "at": "2026-11-15T11:00:00"},
    {"alert_id": "a6", "device": "Old9", "severity": "MAJOR", "type": "x", "message": "보고서에 없는 장비",
     "raw_line": "", "at": "2026-11-15T09:30:00"},
    {"alert_id": "a7", "device": "Agg1", "severity": "MAJOR", "type": "keyword_down", "message": "감시 전",
     "raw_line": "Et9 is down", "at": "2026-11-15T08:00:00", "history": True},
]
OPEN = [{"device": "Core1", "component_id": "mlag:peer-link", "condition": "mlag_peer_down", "since": 1.0,
         "raw": "%MLAG-3-PEER_DOWN", "alert_ids": ["a4"]}]


@pytest.fixture
def round_(root):
    """고객사 ABC 의 회차 2026-11 — 장비 Core1·Agg1 의 점검 로그."""
    from api.customer_profile_api import CustomerProfileApiMixin
    from api.project_core_api import ProjectCoreApiMixin
    from engine import device_inventory as di
    from engine import project_manager as pm

    class Api(CustomerProfileApiMixin, ProjectCoreApiMixin):
        pass

    api = Api()
    api.create_customer("ABC")
    cid = next(c["id"] for c in api.get_customer_profiles())
    api.create_inspection_profile(cid, "2026-11", inspection_date="2026-11-15")
    project = next(p["id"] for c in api.get_customer_profiles() for p in c["profiles"])
    di.save_inventory({"defaults": {}, "devices": [
        {"name": n, "management_ip": f"10.0.0.{i}", "enabled": True} for i, n in enumerate(("Core1", "Agg1"), 1)]},
        pm.project_paths(project)["device_inventory"])
    for name in ("Core1", "Agg1"):
        write_log("ABC", "2026-11", "2026-11-15_090000", name, eos_log(name, uptime="3 days"))
    return project


def save_state(alerts=ALERTS, open_conditions=OPEN):
    from engine import realtime_audit
    from engine import realtime_monitor as rtm
    path = realtime_audit.state_path("ABC", "2026-11")
    assert rtm.save_snapshot(path, {
        "version": 1, "saved_at": time.time(), "devices": ["Core1", "Agg1", "Old9"],
        "alerts": alerts, "open_conditions": open_conditions,
        "observed": {"Core1": {"syslog": 3, "commands": 5}, "Agg1": {"syslog": 0, "commands": 2}}})
    return path


# --------------------------------------------------------------------------- 1·2. 이력 만들기

def test_no_watch_in_this_round_means_no_audit(round_):
    from engine import realtime_audit
    assert realtime_audit.audit("ABC", "2026-11") is None


def test_audit_rows_outcomes_and_counts(round_):
    from engine import realtime_audit
    save_state()
    audit = realtime_audit.audit("ABC", "2026-11", devices=["Core1", "Agg1"])
    rows = audit["rows"]
    assert [r["device"] + ":" + r["outcome"] for r in rows] == [
        "Agg1:open", "Core1:auto", "Agg1:manual", "Core1:ignored", "Core1:open"]       # 오래된 순, 과거 기록 빠짐
    early, auto, manual, ignored, still = rows
    assert early["earlier"] is True and early["at"] == "11-15 08:00:00"
    assert auto["outcome_label"] == "자동 복구" and auto["duration_sec"] == 120.0
    assert auto["note"] == "복구 근거: Core1(config-if-Et1)#no shutdown" and auto["closed_at"] == "11-15 10:02:00"
    assert manual["outcome_label"] == "수동 해결" and manual["note"] == "현장 확인 후 정상"
    assert manual["duration_sec"] is None and manual["closed_at"] == "11-15 10:20:00"
    assert ignored["outcome_label"] == "무시" and ignored["note"] == "계획된 작업"
    assert ignored["closed_at"] == "11-15 10:11:00"
    assert still["outcome_label"] == "미해결" and still["severity_label"] == "심각" and still["closed_at"] == ""
    assert audit["counts"] == {"total": 5, "auto": 1, "manual": 1, "ignored": 1, "open": 2,
                               "CRITICAL": 2, "MAJOR": 2, "WARNING": 1}
    assert audit["excluded"] == 1 and audit["watched"] == 2 and audit["syslog_devices"] == 1
    assert audit["period"] == ("11-15 08:00:00", "11-15 11:00:00")
    [fault] = audit["open_faults"]
    assert fault["device"] == "Core1" and fault["label"] == "MLAG 피어 끊김" and fault["since"]


def test_secrets_are_masked(round_):
    from engine import realtime_audit
    save_state()
    audit = realtime_audit.audit("ABC", "2026-11")
    text = json.dumps(audit["rows"], ensure_ascii=False)
    assert SECRET_HASH not in text and "username admin secret 5" in text
    assert audit["excluded"] == 0 and any(r["device"] == "Old9" for r in audit["rows"])   # 장비를 안 주면 전부


def test_the_saved_state_is_only_read(round_):
    from engine import realtime_audit
    path = Path(save_state())
    save_state(alerts=ALERTS[:2])                                    # 직전 저장본(.bak)이 생긴다
    path.write_text("{ 깨진 저장본", encoding="utf-8")
    audit = realtime_audit.audit("ABC", "2026-11")
    assert audit is not None and audit["counts"]["total"] == 6       # .bak(첫 저장본 — 과거 기록 뺀 6건)을 읽었다
    assert path.read_text(encoding="utf-8") == "{ 깨진 저장본"        # 옮기거나 고치지 않았다
    assert not list(path.parent.glob("realtime_state.json.corrupt-*"))


# --------------------------------------------------------------------------- 3. 보고서

def test_excel_report_gets_the_audit_sheet(round_):
    from openpyxl import load_workbook
    from engine import inspection_report_builder as builder
    save_state()
    result = builder.export("ABC", "2026-11", "excel", project_id=round_)
    workbook = load_workbook(result["path"])
    assert workbook.sheetnames[-1] == "실시간감시이력"
    sheet = workbook["실시간감시이력"]
    assert [c.value for c in sheet[3]] == ["발생", "장비", "등급", "내용", "근거 줄", "처리", "처리 시각", "지속(초)", "메모"]
    values = [[c.value for c in row] for row in sheet.iter_rows(min_row=4, values_only=False)]
    assert [v[5] for v in values] == ["미해결", "자동 복구", "수동 해결", "무시", "미해결"]
    assert "MLAG 피어 끊김" in sheet["A2"].value and "경고 1건 제외" in sheet["A2"].value     # Old9 는 보고서에 없다
    assert SECRET_HASH not in json.dumps(values, ensure_ascii=False, default=str)


def test_excel_report_without_a_watch_has_no_audit_sheet(round_):
    from openpyxl import load_workbook
    from engine import inspection_report_builder as builder
    result = builder.export("ABC", "2026-11", "excel", project_id=round_)
    assert "실시간감시이력" not in load_workbook(result["path"]).sheetnames


def test_pdf_report_gets_the_appendix_and_survives_a_broken_audit(round_, monkeypatch):
    import report.inspection_pdf as pdf
    from engine import inspection_report_builder as builder
    from engine import realtime_audit
    save_state()
    pages = []
    real_page = pdf.audit_page
    monkeypatch.setattr(pdf, "audit_page", lambda audit, site: pages.append(audit) or real_page(audit, site))
    result = builder.export("ABC", "2026-11", "pdf", project_id=round_)
    assert Path(result["path"]).read_bytes()[:4] == b"%PDF" and pages[0]["counts"]["total"] == 5

    def broken(*args, **kwargs):
        raise RuntimeError("이력 실패")

    monkeypatch.setattr(realtime_audit, "audit", broken)
    pages.clear()
    again = builder.export("ABC", "2026-11", "pdf", project_id=round_)
    assert Path(again["path"]).read_bytes()[:4] == b"%PDF" and pages == []          # 부록 없이 보고서는 만든다


def test_audit_page_text():
    from reportlab.platypus import Paragraph, Table
    import report.inspection_pdf as pdf
    pdf.register_fonts()
    audit = {"rows": [], "counts": {"total": 0, "auto": 0, "manual": 0, "ignored": 0, "open": 0,
                                    "CRITICAL": 0, "MAJOR": 0, "WARNING": 0},
             "open_faults": [], "period": None, "watched": 2, "syslog_devices": 0, "excluded": 0}
    story = pdf.audit_page(audit, "ABC")
    texts = [f.getPlainText() for f in story if isinstance(f, Paragraph)]
    assert "부록. ABC 실시간 감시 이력" in texts and "감시하는 동안 경고가 없었습니다." in texts
    assert not any(isinstance(f, Table) for f in story)
    assert [pdf._duration_text(s) for s in (45, 120, 125, 3725)] == ["45초", "2분", "2분 5초", "1시간 2분"]


# --------------------------------------------------------------------------- 4. 보고서 전 저장

def test_reports_flush_the_watch_state_first(monkeypatch):
    from api import inspection_report_api
    from api.inspection_report_api import InspectionReportApiMixin
    from api.realtime_watch_api import RealtimeWatchApiMixin, flush_realtime_state
    order = []

    class Service:
        def flush_state(self):
            order.append("flush")
            return True

    class Stub(InspectionReportApiMixin, RealtimeWatchApiMixin):
        def _project(self):
            return "proj"

        def _resolve_active_customer_profile_names(self):
            return ("ABC", "2026-11")

    bare = Stub()
    assert flush_realtime_state(bare) is False and "_realtime_service" not in bare.__dict__
    monkeypatch.setattr(inspection_report_api.builder, "export",
                        lambda *a, **k: order.append("export") or {"path": "x.xlsx"})
    assert bare.export_inspection_report()["ok"] and order == ["export"]   # 감시를 쓰지 않았다 — 저장할 것 없음
    order.clear()
    watched = Stub()
    watched._realtime_service = Service()
    assert watched.export_inspection_report()["ok"] and order == ["flush", "export"]
