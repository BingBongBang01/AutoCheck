"""자격증명 금고 — IMPROVEMENT_PLAN 4-1.

장비 비밀번호·붙여넣은 개인키·passphrase·기본 비밀번호와 클라우드 AI API 키가 디스크에 평문으로 있었고, 장비
목록 화면은 get_devices 로 비밀번호·개인키 평문을 받아 input value 로 DOM 에 넣었다. 이제:
  1. 디스크에는 Windows DPAPI 로 봉인한 'vault:dpapi:v1:…' 문자열로 둔다(core/secret_vault).
  2. 화면에는 저장 여부(has_*)만 보내고, 화면이 보낸 목록은 장비 id 로 기존 비밀값을 잇는다.
  3. 엑셀 내보내기는 비밀번호를 기본으로 뺀다.
  4. 예전 평문은 앱 시작 때 한 번 봉인한다(업그레이드 1). 예전 credential.json 사본은 더 만들지 않는다.
"""
import json
import sys
import tempfile
from pathlib import Path

import pytest
import yaml

from core import secret_vault as vault
from core.paths import AppPaths

pytestmark = pytest.mark.skipif(sys.platform != "win32" or not vault.available(),
                                reason="Windows DPAPI 가 있어야 봉인을 확인할 수 있다")

UNREADABLE = vault.PREFIX + "AAAA"          # 형식은 맞지만 이 PC 에서 풀 수 없는 봉인 값(다른 PC 에서 온 것처럼)


@pytest.fixture
def root():
    original = AppPaths._user_data_root
    AppPaths._user_data_root = Path(tempfile.mkdtemp()) / "AutoCheck"
    AppPaths.forget_ensured()
    yield AppPaths._user_data_root
    AppPaths._user_data_root = original
    AppPaths.forget_ensured()


# --------------------------------------------------------------------------- 1. 금고

def test_seal_round_trip_and_passthrough():
    sealed = vault.seal('p@ss "한글" &<>')
    assert vault.is_sealed(sealed) and 'p@ss' not in sealed
    assert vault.reveal(sealed) == 'p@ss "한글" &<>'
    assert vault.seal(sealed) == sealed                  # 두 번 봉인하지 않는다
    assert vault.seal("") == "" and vault.seal(None) is None
    assert vault.reveal("평문") == "평문"                  # 이전 전 값은 그대로 읽힌다


def test_unreadable_values_raise_and_fields_report_them():
    with pytest.raises(vault.VaultError):
        vault.reveal(UNREADABLE)
    data = {"devices": [{"password": vault.seal("a")}, {"password": UNREADABLE}], "api_key": "plain"}
    revealed, failed = vault.reveal_fields(data, {"password", "api_key"})
    assert [d["password"] for d in revealed["devices"]] == ["a", ""] and failed == ["password"]
    sealed, count = vault.seal_fields(revealed, {"password", "api_key"})
    assert count == 2 and vault.is_sealed(sealed["api_key"])          # 빈 값은 세지 않는다


# --------------------------------------------------------------------------- 2. 장비 목록의 저장 형태

def _inventory(tmp_path, devices, defaults=None):
    path = tmp_path / "device_inventory.yaml"
    path.write_text(yaml.safe_dump({"defaults": defaults or {}, "devices": devices}, allow_unicode=True),
                    encoding="utf-8")
    return str(path)


def test_inventory_is_sealed_on_disk_and_plain_in_memory(tmp_path):
    from engine import device_inventory as di
    # 값마다 '-' 를 넣는다 — base64 에는 나올 수 없는 글자라, 봉인된 문자열에 우연히 같은 글자가 섞여 실패하지 않는다
    # (예전엔 기본 비밀번호가 'dp' 라 봉인값 안의 'dp' 에 걸려 가끔 실패했다).
    path = _inventory(tmp_path, [{"name": "Core1", "password": "pw-1", "key_content": "-----BEGIN KEY-----",
                                  "key_passphrase": "pp-1"}], {"default_password": "dp-1"})
    inv = di.load_inventory(path)                                     # 이전 전(평문) 파일도 읽힌다
    assert inv["devices"][0]["password"] == "pw-1" and inv["defaults"]["default_password"] == "dp-1"
    di.save_inventory(inv, path)
    text = Path(path).read_text(encoding="utf-8")
    assert [value for value in ("pw-1", "BEGIN KEY", "pp-1", "dp-1") if value in text] == []
    assert text.count(vault.PREFIX) == 4
    again = di.load_inventory(path)
    assert again["devices"][0]["password"] == "pw-1" and again["devices"][0]["key_passphrase"] == "pp-1"
    assert again["devices"][0]["id"] == inv["devices"][0]["id"] != ""   # id 가 파일에 남는다
    assert inv["devices"][0]["password"] == "pw-1"                    # 저장해도 메모리 값은 평문 그대로


def test_values_from_another_pc_are_kept_and_reported(tmp_path):
    from engine import device_inventory as di
    path = _inventory(tmp_path, [{"name": "Core1", "password": UNREADABLE}])
    inv = di.load_inventory(path)
    device = inv["devices"][0]
    assert device["password"] == ""                                  # 접속하는 쪽에는 '없음'
    assert di.public_device(device)["unreadable_secrets"] == ["password"]
    di.save_inventory(inv, path)                                      # 다른 칸을 고쳐 저장해도
    assert UNREADABLE in Path(path).read_text(encoding="utf-8")       # 그 값을 지우지 않는다


# --------------------------------------------------------------------------- 3. 화면 API

@pytest.fixture
def api(tmp_path):
    from api.inventory_api import InventoryApiMixin
    from engine import device_inventory as di
    paths = {"device_inventory": str(tmp_path / "labs" / "p1" / "device_inventory.yaml"),
             "lab_meta": str(tmp_path / "lab_meta.yaml"), "ip_allocation": str(tmp_path / "ip_allocation.yaml")}

    class Stub(InventoryApiMixin):
        def _paths(self):
            return paths

        def _project(self):
            return "p1"

        def _load_inventory(self, p):
            return di.load_inventory(p["device_inventory"], p["lab_meta"], p["ip_allocation"])

        def _resolve_active_customer_profile_names(self):
            return "고객사A", "2026-11"

    stub = Stub()
    di.save_inventory({"defaults": {"default_username": "ops", "default_password": "dpw"},
                       "devices": [{"name": "Core1", "management_ip": "10.0.0.1", "password": "pw1",
                                    "key_content": "KEY", "key_passphrase": "pp", "enabled": True}]},
                      paths["device_inventory"])
    stub.paths = paths
    return stub


def _stored(api):
    return api._load_inventory(api.paths)


def test_the_screen_never_receives_secrets(api):
    device = api.get_devices()[0]
    assert not {"password", "key_content", "key_passphrase"} & set(device)
    assert device["has_password"] and device["has_key_content"] and device["has_key_passphrase"]
    defaults = api.get_inventory_defaults()
    assert "default_password" not in defaults and defaults["has_default_password"]


def test_saving_without_secrets_keeps_them_even_after_a_rename(api):
    device = api.get_devices()[0]
    assert api.save_devices([dict(device, name="Core1-renamed", management_ip="10.0.0.9")]) is True
    stored = _stored(api)["devices"][0]
    assert (stored["name"], stored["password"], stored["key_content"], stored["key_passphrase"]) == \
        ("Core1-renamed", "pw1", "KEY", "pp")


def test_new_values_replace_and_clear_secrets_removes(api):
    device = api.get_devices()[0]
    api.save_devices([dict(device, password="new-pw", clear_secrets=["key_passphrase"])])
    stored = _stored(api)["devices"][0]
    assert stored["password"] == "new-pw" and stored["key_passphrase"] == "" and stored["key_content"] == "KEY"


def test_a_new_row_keeps_its_password_across_later_saves(api):
    """화면이 첫 저장 뒤 비밀번호를 잊어도, 다음 자동 저장이 같은 장비로 이어진다(화면이 붙인 id)."""
    rows = api.get_devices() + [{"id": "client-1", "name": "Agg1", "password": "agg-pw"}]
    api.save_devices(rows)
    later = [dict(d) for d in api.get_devices()]
    later[1]["management_ip"] = "10.0.0.2"
    api.save_devices(later)
    assert _stored(api)["devices"][1]["password"] == "agg-pw"


def test_rows_without_an_id_fall_back_to_the_name(api):
    device = {k: v for k, v in api.get_devices()[0].items() if k != "id"}
    api.save_devices([device])
    assert _stored(api)["devices"][0]["password"] == "pw1"


def test_defaults_keep_the_password_unless_changed_or_cleared(api):
    api.save_inventory_defaults({"auto_save": False})
    assert _stored(api)["defaults"]["default_password"] == "dpw"
    api.save_inventory_defaults({"default_username": "ops2", "default_password": "dpw2"})
    assert _stored(api)["defaults"]["default_password"] == "dpw2"
    api.save_inventory_defaults({"clear_secrets": ["default_password"]})
    assert _stored(api)["defaults"]["default_password"] == ""


def test_connection_check_uses_the_stored_secret(api, monkeypatch):
    from engine import device_probe
    seen = []
    monkeypatch.setattr(device_probe, "probe_device", lambda device, defaults, **kw: seen.append(device) or {"ok": True})
    api.probe_device_config(dict(api.get_devices()[0], management_ip="10.0.0.5"))   # 화면엔 비밀번호가 없다
    assert seen[0]["password"] == "pw1" and seen[0]["management_ip"] == "10.0.0.5"


def test_excel_export_leaves_passwords_out_by_default(api, tmp_path, monkeypatch):
    import openpyxl
    from api import window_ref
    target = tmp_path / "export.xlsx"

    class Window:
        def create_file_dialog(self, *args, **kwargs):
            return str(target)

    monkeypatch.setattr(window_ref, "get_window", lambda: Window())
    result = api.export_devices()
    assert result["path"] == str(target) and "notice" in result
    cells = [c.value for row in openpyxl.load_workbook(target).active.iter_rows() for c in row]
    assert "pw1" not in cells and "KEY" not in cells
    result = api.export_devices(True)
    assert "warning" in result
    assert "pw1" in [c.value for row in openpyxl.load_workbook(target).active.iter_rows() for c in row]


# --------------------------------------------------------------------------- 4. AI API 키

@pytest.fixture
def ai(root):
    from api.cloud_ai_api import CloudAiApiMixin
    from api.local_ai_api import LocalAiApiMixin

    class Stub(LocalAiApiMixin, CloudAiApiMixin):
        pass

    return Stub()


def test_api_keys_are_sealed_on_disk(ai):
    ai.add_cloud_api("main", "anthropic", api_key="sk-ant-secret-0123456789abcdef")
    text = Path(ai.AI_CONFIG_PATH).read_text(encoding="utf-8")
    assert "sk-ant" not in text and vault.PREFIX in text
    entry = ai._load_ai_config()["providers"][0]["entries"][0]
    assert entry["api_key"] == "sk-ant-secret-0123456789abcdef"
    assert ai.get_cloud_apis()["entries"][0]["has_key"] is True


def test_an_unreadable_api_key_is_reported_and_kept(ai):
    Path(ai.AI_CONFIG_PATH).parent.mkdir(parents=True, exist_ok=True)
    Path(ai.AI_CONFIG_PATH).write_text(yaml.safe_dump({"providers": [{"type": "cloud_apis", "entries": [
        {"id": "e1", "name": "old", "provider": "anthropic", "enabled": True, "api_key": UNREADABLE}]}]}),
        encoding="utf-8")
    listed = ai.get_cloud_apis()["entries"][0]
    assert listed["has_key"] is False and listed["key_unreadable"] is True
    ai._save_ai_config(ai._load_ai_config())
    assert UNREADABLE in Path(ai.AI_CONFIG_PATH).read_text(encoding="utf-8")


# --------------------------------------------------------------------------- 5. 예전 평문의 1회 이전

def test_upgrade_seals_every_known_secret_file_once(root):
    from engine import migration_manager
    files = {
        root / "labs" / "p1" / "device_inventory.yaml":
            ("yaml", {"defaults": {"default_password": "dp"}, "devices": [{"name": "a", "password": "pw"}]}),
        root / "labs" / "p1" / "ip_allocation.yaml":
            ("yaml", {"default_credentials": {"username": "u", "password": "lp"}}),
        root / "data" / "C" / "P" / "profile" / "credential.json": ("json", {"username": "u", "password": "cp"}),
        root / "ai_config.yaml": ("yaml", {"providers": [{"type": "cloud_apis", "entries": [{"api_key": "k1"}]}]}),
    }
    for path, (kind, data) in files.items():
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(data) if kind == "json" else yaml.safe_dump(data), encoding="utf-8")
    exported = root / "data" / "C" / "P" / "inventory" / "devices.csv"
    exported.parent.mkdir(parents=True)
    exported.write_text("name,password\na,pw\n", encoding="utf-8")       # 사용자가 내보낸 파일은 건드리지 않는다

    assert migration_manager.upgrade_1_seal_secrets(None) == {"files": 4, "values": 5}
    for path in files:
        text = path.read_text(encoding="utf-8")
        assert vault.PREFIX in text and not any(f": {v}" in text or f'"{v}"' in text
                                               for v in ("pw", "dp", "lp", "cp", "k1"))
    assert exported.read_text(encoding="utf-8") == "name,password\na,pw\n"
    assert migration_manager.upgrade_1_seal_secrets(None) == {"files": 0, "values": 0}   # 멱등


def test_legacy_inventory_migration_reads_sealed_credentials(tmp_path):
    from engine import device_inventory as di
    lab_meta, ip_alloc = tmp_path / "lab_meta.yaml", tmp_path / "ip_allocation.yaml"
    lab_meta.write_text(yaml.safe_dump({"devices": [{"name": "Core1"}]}), encoding="utf-8")
    ip_alloc.write_text(yaml.safe_dump({"allocations": [{"node_name": "Core1", "ip": "10.0.0.1",
                                                         "password": vault.seal("lp")}],
                                        "default_credentials": {"username": "u", "password": vault.seal("dp")}}),
                        encoding="utf-8")
    inv = di.load_inventory(str(tmp_path / "device_inventory.yaml"), str(lab_meta), str(ip_alloc))
    assert inv["devices"][0]["password"] == "lp" and inv["defaults"]["default_password"] == "dp"


def test_project_zip_reveals_only_when_credentials_are_included(tmp_path):
    from api.project_core_api import _revealed_yaml, _scrubbed_yaml
    from engine.secret_files import seal_secret_files
    path = tmp_path / "device_inventory.yaml"
    path.write_text(yaml.safe_dump({"devices": [{"name": "a", "password": vault.seal("pw")}]}), encoding="utf-8")
    assert "pw" not in _scrubbed_yaml(str(path)).replace("password", "")      # 기본: 비운다
    assert yaml.safe_load(_revealed_yaml(str(path)))["devices"][0]["password"] == "pw"   # 포함: 다른 PC 에서 쓰게 푼다
    imported = tmp_path / "imported" / "device_inventory.yaml"                # 불러오면 바로 봉인한다
    imported.parent.mkdir()
    imported.write_text(_revealed_yaml(str(path)), encoding="utf-8")
    assert seal_secret_files(tmp_path / "imported") == {"files": 1, "values": 1}
    assert vault.PREFIX in imported.read_text(encoding="utf-8")


def test_no_more_credential_json_copies(root):
    from engine import profile_manager
    assert not hasattr(profile_manager.ProfileManager, "migrate_credentials_from_yaml")
